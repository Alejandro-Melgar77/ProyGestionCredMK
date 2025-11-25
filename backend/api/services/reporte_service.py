import speech_recognition as sr
import json
import re
from datetime import datetime, timedelta
from dateutil.parser import parse
import os
import io

from django.http import HttpResponse
from django.db.models import Count, Sum, Avg
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
import openpyxl
from openpyxl.styles import Font, Alignment
from pydub import AudioSegment  # 🔥 Para convertir audio a WAV

from ..models import Cliente, SolicitudCredito, TransaccionPago

# -------------------------------------------------------------------------
#  PROCESADOR DE COMANDOS POR VOZ
# -------------------------------------------------------------------------

class VoiceCommandProcessor:
    def __init__(self):
        self.keywords = {
            'tipo_reporte': {
                'creditos': ['créditos', 'préstamos', 'solicitudes'],
                'clientes': ['clientes', 'usuarios', 'personas'],
                'pagos': ['pagos', 'cuotas', 'transacciones'],
                'riesgo': ['riesgo', 'evaluación', 'scoring']
            },
            'periodo': {
                'hoy': ['hoy', 'hoy día', 'este día'],
                'semana': ['semana', 'última semana', 'esta semana'],
                'mes': ['mes', 'este mes', 'último mes'],
                'año': ['año', 'este año', 'último año']
            },
            'estado': {
                'aprobado': ['aprobado', 'aceptado', 'autorizado'],
                'rechazado': ['rechazado', 'denegado', 'negado'],
                'pendiente': ['pendiente', 'en proceso', 'en revisión'],
                'desembolsado': ['desembolsado', 'pagado', 'entregado']
            },
            'producto': {
                'consumo': ['consumo', 'personal', 'individual'],
                'vivienda': ['vivienda', 'casa', 'hipotecario'],
                'pyme': ['pyme', 'empresa', 'negocio'],
                'vehiculo': ['vehículo', 'auto', 'carro']
            }
        }

    def transcribe_audio(self, audio_file):
        """
        Convierte audio a WAV con pydub y luego transcribe usando SpeechRecognition.
        Compatible con .webm subido desde frontend.
        """
        try:
            # Convertir a WAV
            audio = AudioSegment.from_file(audio_file, format="webm")
            wav_buffer = io.BytesIO()
            audio.export(wav_buffer, format="wav")
            wav_buffer.seek(0)

            # Reconocimiento de voz
            r = sr.Recognizer()
            with sr.AudioFile(wav_buffer) as source:
                r.adjust_for_ambient_noise(source)
                audio_data = r.record(source)
                text = r.recognize_google(audio_data, language="es-ES")
            return text.lower()

        except sr.UnknownValueError:
            raise Exception("No se pudo entender el audio")
        except sr.RequestError as e:
            raise Exception(f"Error en reconocimiento: {e}")

    def extract_filters_from_text(self, text):
        filters = {
            'tipo_reporte': 'creditos',
            'fecha_inicio': None,
            'fecha_fin': None,
            'estado': None,
            'tipo_producto': None,
            'formato': 'pdf'
        }

        # Tipo de reporte
        for report_type, keywords in self.keywords['tipo_reporte'].items():
            if any(k in text for k in keywords):
                filters['tipo_reporte'] = report_type
                break

        # Rango de fechas
        filters.update(self._extract_dates(text))

        # Estado
        for estado, keywords in self.keywords['estado'].items():
            if any(k in text for k in keywords):
                filters['estado'] = estado
                break

        # Producto
        for producto, keywords in self.keywords['producto'].items():
            if any(k in text for k in keywords):
                filters['tipo_producto'] = producto
                break

        # Formato
        if 'excel' in text or 'hoja de cálculo' in text:
            filters['formato'] = 'excel'
        elif 'texto' in text:
            filters['formato'] = 'texto'

        return filters

    def _extract_dates(self, text):
        dates = {'fecha_inicio': None, 'fecha_fin': None}
        today = datetime.now()

        if any(w in text for w in self.keywords['periodo']['hoy']):
            dates['fecha_inicio'] = dates['fecha_fin'] = today.strftime('%Y-%m-%d')
        elif any(w in text for w in self.keywords['periodo']['semana']):
            dates['fecha_inicio'] = (today - timedelta(days=7)).strftime('%Y-%m-%d')
            dates['fecha_fin'] = today.strftime('%Y-%m-%d')
        elif any(w in text for w in self.keywords['periodo']['mes']):
            dates['fecha_inicio'] = (today - timedelta(days=30)).strftime('%Y-%m-%d')
            dates['fecha_fin'] = today.strftime('%Y-%m-%d')
        elif any(w in text for w in self.keywords['periodo']['año']):
            dates['fecha_inicio'] = (today - timedelta(days=365)).strftime('%Y-%m-%d')
            dates['fecha_fin'] = today.strftime('%Y-%m-%d')

        return dates

    def process_voice_command(self, audio_file):
        try:
            transcribed_text = self.transcribe_audio(audio_file)
            filters = self.extract_filters_from_text(transcribed_text)
            return {'transcribed_text': transcribed_text, 'filters': filters, 'success': True}
        except Exception as e:
            return {'success': False, 'error': str(e), 'filters': {}}

# -------------------------------------------------------------------------
#  GENERADOR DE REPORTES
# -------------------------------------------------------------------------

class ReporteGenerator:

    # ------------------ CRÉDITOS ------------------
    @staticmethod
    def generar_reporte_creditos(filtros, formato):
        creditos = SolicitudCredito.objects.all()
        if filtros.get('fecha_inicio'):
            creditos = creditos.filter(created_at__gte=filtros['fecha_inicio'])
        if filtros.get('fecha_fin'):
            creditos = creditos.filter(created_at__lte=filtros['fecha_fin'])
        if filtros.get('estado'):
            creditos = creditos.filter(estado=filtros['estado'])
        if filtros.get('tipo_producto'):
            creditos = creditos.filter(tipo_credito=filtros['tipo_producto'])

        datos = creditos.select_related('cliente')

        if formato == 'pdf':
            return ReporteGenerator._pdf_creditos(datos)
        elif formato == 'excel':
            return ReporteGenerator._excel_creditos(datos)
        else:
            return ReporteGenerator._texto_creditos(datos)

    @staticmethod
    def _pdf_creditos(datos):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()

        elements = [Paragraph("Reporte de Créditos", styles['Heading1'])]
        data = [['Cliente', 'Monto', 'Estado', 'Fecha', 'Producto']]

        for c in datos:
            data.append([
                c.cliente.user.get_full_name(),
                f"{c.monto:,.2f}",
                c.estado,
                c.created_at.strftime('%d/%m/%Y'),
                c.tipo_credito
            ])

        table = Table(data)
        table.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 1, colors.black)]))
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return buffer, "reporte_creditos.pdf", "application/pdf"

    @staticmethod
    def _excel_creditos(datos):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Créditos"

        headers = ['Cliente', 'Monto', 'Estado', 'Fecha', 'Producto']
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h).font = Font(bold=True)

        for row, c in enumerate(datos, 2):
            ws.cell(row=row, column=1, value=c.cliente.user.get_full_name())
            ws.cell(row=row, column=2, value=float(c.monto))
            ws.cell(row=row, column=3, value=c.estado)
            ws.cell(row=row, column=4, value=c.created_at.strftime('%d/%m/%Y'))
            ws.cell(row=row, column=5, value=c.tipo_credito)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer, "reporte_creditos.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @staticmethod
    def _texto_creditos(datos):
        contenido = "REPORTE DE CRÉDITOS\n\n"
        for c in datos:
            contenido += (
                f"Cliente: {c.cliente.user.get_full_name()}\n"
                f"Monto: {c.monto:,.2f}\n"
                f"Estado: {c.estado}\n"
                f"Fecha: {c.created_at.strftime('%d/%m/%Y')}\n"
                f"Producto: {c.tipo_credito}\n"
                "-----------------------------------------\n"
            )
        buffer = io.BytesIO(contenido.encode())
        return buffer, "reporte_creditos.txt", "text/plain"

    # --------------------------------------------------------------------------------
    # CLIENTES
    # --------------------------------------------------------------------------------
    @staticmethod
    def generar_reporte_clientes(filtros, formato):
        clientes = Cliente.objects.all()
        if filtros.get('fecha_inicio'):
            clientes = clientes.filter(fecha_registro__gte=filtros['fecha_inicio'])
        if filtros.get('fecha_fin'):
            clientes = clientes.filter(fecha_registro__lte=filtros['fecha_fin'])

        if formato == 'pdf':
            return ReporteGenerator._pdf_clientes(clientes)
        elif formato == 'excel':
            return ReporteGenerator._excel_clientes(clientes)
        else:
            return ReporteGenerator._texto_clientes(clientes)

    @staticmethod
    def _pdf_clientes(datos):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = [Paragraph("Reporte de Clientes", styles["Heading1"])]
        data = [["Nombre", "Email", "Fecha Registro"]]
        for c in datos:
            data.append([c.user.get_full_name(), c.user.email, c.fecha_registro.strftime('%d/%m/%Y')])
        table = Table(data)
        table.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 1, colors.black)]))
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return buffer, "reporte_clientes.pdf", "application/pdf"

    @staticmethod
    def _excel_clientes(datos):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clientes"
        headers = ["Nombre", "Email", "Fecha Registro"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h).font = Font(bold=True)
        for row, c in enumerate(datos, 2):
            ws.cell(row=row, column=1, value=c.user.get_full_name())
            ws.cell(row=row, column=2, value=c.user.email)
            ws.cell(row=row, column=3, value=c.fecha_registro.strftime('%d/%m/%Y'))
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer, "reporte_clientes.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @staticmethod
    def _texto_clientes(datos):
        contenido = "REPORTE DE CLIENTES\n\n"
        for c in datos:
            contenido += (
                f"Nombre: {c.user.get_full_name()}\n"
                f"Email: {c.user.email}\n"
                f"Fecha Registro: {c.fecha_registro.strftime('%d/%m/%Y')}\n"
                "-----------------------------------------\n"
            )
        buffer = io.BytesIO(contenido.encode())
        return buffer, "reporte_clientes.txt", "text/plain"

    # --------------------------------------------------------------------------------
    # PAGOS
    # --------------------------------------------------------------------------------
    @staticmethod
    def generar_reporte_pagos(filtros, formato):
        pagos = TransaccionPago.objects.select_related('cuota__plan', 'cuota__plan__solicitud', 'cuota__plan__solicitud__cliente')
        if filtros.get('fecha_inicio'):
            pagos = pagos.filter(fecha_transaccion__gte=filtros['fecha_inicio'])
        if filtros.get('fecha_fin'):
            pagos = pagos.filter(fecha_transaccion__lte=filtros['fecha_fin'])
        if filtros.get('estado'):
            pagos = pagos.filter(estado=filtros['estado'])

        if formato == 'pdf':
            return ReporteGenerator._pdf_pagos(pagos)
        elif formato == 'excel':
            return ReporteGenerator._excel_pagos(pagos)
        else:
            return ReporteGenerator._texto_pagos(pagos)

    @staticmethod
    def _pdf_pagos(datos):
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = [Paragraph("Reporte de Pagos", styles["Heading1"])]
        data = [["Cliente", "Crédito", "Monto", "Estado", "Fecha"]]

        for p in datos:
            solicitud = p.cuota.plan.solicitud
            data.append([
                solicitud.cliente.user.get_full_name(),
                str(solicitud.id),
                f"{p.monto:,.2f}",
                p.estado,
                p.fecha_transaccion.strftime('%d/%m/%Y')
            ])

        table = Table(data)
        table.setStyle(TableStyle([('GRID', (0,0), (-1,-1), 1, colors.black)]))
        elements.append(table)
        doc.build(elements)
        buffer.seek(0)
        return buffer, "reporte_pagos.pdf", "application/pdf"

    @staticmethod
    def _excel_pagos(datos):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pagos"
        headers = ["Cliente", "Crédito", "Monto", "Estado", "Fecha"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h).font = Font(bold=True)
        for row, p in enumerate(datos, 2):
            solicitud = p.cuota.plan.solicitud
            ws.cell(row=row, column=1, value=solicitud.cliente.user.get_full_name())
            ws.cell(row=row, column=2, value=str(solicitud.id))
            ws.cell(row=row, column=3, value=float(p.monto))
            ws.cell(row=row, column=4, value=p.estado)
            ws.cell(row=row, column=5, value=p.fecha_transaccion.strftime('%d/%m/%Y'))
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer, "reporte_pagos.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @staticmethod
    def _texto_pagos(datos):
        contenido = "REPORTE DE PAGOS\n\n"
        for p in datos:
            solicitud = p.cuota.plan.solicitud
            contenido += (
                f"Cliente: {solicitud.cliente.user.get_full_name()}\n"
                f"Crédito: {solicitud.id}\n"
                f"Monto: {p.monto:,.2f}\n"
                f"Estado: {p.estado}\n"
                f"Fecha Pago: {p.fecha_transaccion.strftime('%d/%m/%Y')}\n"
                "-----------------------------------------\n"
            )
        buffer = io.BytesIO(contenido.encode())
        return buffer, "reporte_pagos.txt", "text/plain"
